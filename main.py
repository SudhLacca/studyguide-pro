from fastapi import FastAPI, Depends, HTTPException, status, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm
from sqlalchemy.orm import Session
from pydantic import BaseModel
from typing import Optional, List
import os
import shutil
from datetime import date, timedelta # <-- Added for Study Streak logic

from docx import Document
from docx.shared import Pt # <-- Added for professional DOCX spacing
from fastapi.responses import FileResponse
import tempfile
from src.agents.notes_generator import NotesGenerator 

# Import custom database and auth modules
from src.database import models, database
from src.auth import auth
from config import Config

# Import AI and Data Processing Services
from src.data_processing.document_parser import DocumentParser
from src.data_processing.text_chunker import TextChunker
from src.data_processing.embedding_generator import EmbeddingGenerator
from src.vector_store.chroma_store import ChromaStore
from src.gemini.gemini_service import GeminiService

# Import Agents
from src.agents.concept_simplifier import ConceptSimplifier
from src.agents.auto_glossary import AutoGlossaryExtractor
from src.agents.visual_concept_mapper import VisualConceptMapper
from src.agents.question_bank_generator import QuestionBankGenerator

# Initialize the FastAPI app
app = FastAPI(title="StudyGuide Pro API")

# Configure CORS so our React frontend (port 5173) can talk to this backend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173"],  # React Vite default port
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Create the database tables
models.Base.metadata.create_all(bind=database.engine)

# Security Scheme (Updated to match the route)
oauth2_scheme = OAuth2PasswordBearer(tokenUrl="login")

# --- PYDANTIC SCHEMAS (Data Validation) ---
class UserCreate(BaseModel):
    email: str
    password: str

class Token(BaseModel):
    access_token: str
    token_type: str

class SummaryRequest(BaseModel):
    document_name: str
    focus_area: Optional[str] = None

class FlashcardRequest(BaseModel):
    document_name: str
    num_cards: int = 5
    topic: Optional[str] = None

class StudyGuideRequest(BaseModel):
    document_name: str

class SimplifyRequest(BaseModel):
    document_name: str
    text_to_simplify: str
    target_audience: str

class NotesRequest(BaseModel):
    document_name: str

class DocxExportRequest(BaseModel):
    notes_content: str

class QuestionBankRequest(BaseModel):
    document_name: str
    topic: Optional[str] = None


# --- DEPENDENCIES ---
def get_current_user(token: str = Depends(oauth2_scheme), db: Session = Depends(database.get_db)):
    """Verifies the JWT token and returns the current logged-in user"""
    credentials_exception = HTTPException(
        status_code=status.HTTP_401_UNAUTHORIZED,
        detail="Could not validate credentials",
        headers={"WWW-Authenticate": "Bearer"},
    )
    try:
        # Decode the token
        payload = auth.jwt.decode(token, auth.SECRET_KEY, algorithms=[auth.ALGORITHM])
        email: str = payload.get("sub")
        if email is None:
            raise credentials_exception
    except auth.jwt.JWTError:
        raise credentials_exception
        
    user = db.query(models.User).filter(models.User.email == email).first()
    if user is None:
        raise credentials_exception
    return user


# --- AUTHENTICATION ENDPOINTS ---

@app.post("/register", response_model=Token)
def register_user(user: UserCreate, db: Session = Depends(database.get_db)):
    # Check if user already exists
    db_user = db.query(models.User).filter(models.User.email == user.email).first()
    if db_user:
        raise HTTPException(status_code=400, detail="Email already registered")
    
    # Hash password and save user
    hashed_password = auth.get_password_hash(user.password)
    new_user = models.User(email=user.email, hashed_password=hashed_password)
    db.add(new_user)
    db.commit()
    db.refresh(new_user)
    
    # Generate login token instantly
    access_token = auth.create_access_token(data={"sub": new_user.email})
    return {"access_token": access_token, "token_type": "bearer"}

@app.post("/login", response_model=Token)
def login(form_data: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(database.get_db)):
    user = db.query(models.User).filter(models.User.email == form_data.username).first()
    if not user or not auth.verify_password(form_data.password, user.hashed_password):
        raise HTTPException(status_code=400, detail="Incorrect email or password")
    
    # --- STREAK CALCULATION LOGIC ---
    today = date.today()
    if user.last_login_date == today:
        # They already logged in today, do nothing to the streak
        pass 
    elif user.last_login_date == today - timedelta(days=1):
        # They logged in yesterday! Streak continues!
        user.current_streak += 1
    else:
        # They missed a day, or it's their very first login. Reset to 1.
        user.current_streak = 1
        
    # Update the last login date to right now and save
    user.last_login_date = today
    db.commit()
    # --------------------------------
    
    access_token = auth.create_access_token(data={"sub": user.email})
    return {"access_token": access_token, "token_type": "bearer"}


# --- APPLICATION ENDPOINTS (Protected by get_current_user) ---

@app.post("/upload")
async def upload_document(
    file: UploadFile = File(...), 
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(database.get_db) # <-- Added DB Session
):
    """Handles file uploads, parses text, and stores semantic embeddings"""
    file_path = os.path.join(Config.UPLOAD_FOLDER, file.filename)
    
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
        
    try:
        # 1. Parse Document
        parser = DocumentParser()
        text = parser.parse(file_path)
        
        # 2. Chunk Text
        chunker = TextChunker()
        chunks = chunker.chunk_text(text, file.filename)
        
        # 3. Embed & Store
        embedder = EmbeddingGenerator()
        embedded_chunks = embedder.generate_embeddings(chunks)
        
        vector_store = ChromaStore()
        vector_store.add_documents(embedded_chunks, file.filename)
        
        # --- NEW: Update Database Stats ---
        current_user.documents_analyzed += 1
        db.commit()
        
        return {"status": "success", "filename": file.filename, "chunks_processed": len(chunks)}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/summary")
def generate_summary(
    request: SummaryRequest, 
    current_user: models.User = Depends(get_current_user)
):
    vector_store = ChromaStore()
    gemini = GeminiService()
    
    query = f"Summarize the document{' focusing on ' + request.focus_area if request.focus_area else ''}"
    relevant_chunks = vector_store.search(query, request.document_name, n_results=5)
    
    prompt = f"Create a comprehensive summary of the following content{' with focus on ' + request.focus_area if request.focus_area else ''}. Use plain text only:\n\n"
    summary = gemini.generate_with_context(prompt, relevant_chunks)
    
    return {"summary": summary}

@app.post("/flashcards")
def generate_flashcards(
    request: FlashcardRequest, 
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(database.get_db) # <-- Added DB Session
):
    vector_store = ChromaStore()
    gemini = GeminiService()
    
    query = f"Generate flashcards{' about ' + request.topic if request.topic else ''}"
    relevant_chunks = vector_store.search(query, request.document_name, n_results=5)
    
    prompt = f"""Create {request.num_cards} flashcards{' focusing on ' + request.topic if request.topic else ''}.
    Format EXACTLY as:
    Q: [Question]
    A: [Answer]
    One blank line between cards."""
    
    flashcards_text = gemini.generate_with_context(prompt, relevant_chunks)
    
    # Parse the text into a list of dictionaries for the frontend
    flashcards = []
    lines = flashcards_text.strip().split('\n')
    current_q, current_a = "", ""
    for line in lines:
        line = line.strip()
        if line.startswith('Q:'):
            if current_q and current_a:
                flashcards.append({'question': current_q, 'answer': current_a})
            current_q = line[2:].strip()
            current_a = ""
        elif line.startswith('A:'):
            current_a = line[2:].strip()
    if current_q and current_a:
        flashcards.append({'question': current_q, 'answer': current_a})
        
    # --- NEW: Update Database Stats ---
    if len(flashcards) > 0:
        current_user.flashcards_generated += len(flashcards)
        db.commit()
        
    return {"flashcards": flashcards}

@app.post("/study-guide")
def generate_study_guide(
    request: StudyGuideRequest, 
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(database.get_db) # <-- Added DB Session
):
    mapper = VisualConceptMapper()
    glossary_extractor = AutoGlossaryExtractor()
    
    mindmap = mapper.generate_mindmap(request.document_name)
    glossary = glossary_extractor.generate_glossary(request.document_name)
    
    # --- NEW: Update Database Stats ---
    current_user.study_guides_created += 1
    db.commit()
    
    return {
        "mindmap_code": mindmap,
        "glossary": glossary
    }

@app.post("/simplify")
def simplify_concept(
    request: SimplifyRequest, 
    current_user: models.User = Depends(get_current_user)
):
    simplifier = ConceptSimplifier()
    simplified_text = simplifier.simplify_text(
        request.document_name, 
        request.text_to_simplify, 
        request.target_audience
    )
    return {"simplified_text": simplified_text}

@app.post("/generate-notes")
def generate_detailed_notes(
    request: NotesRequest, 
    current_user: models.User = Depends(get_current_user)
):
    """Generates detailed markdown notes using AI"""
    agent = NotesGenerator()
    notes = agent.generate_notes(request.document_name)
    return {"notes": notes}

@app.post("/export-docx")
def export_to_docx(
    request: DocxExportRequest, 
    current_user: models.User = Depends(get_current_user)
):
    """Converts markdown notes into a well-spaced, downloadable DOCX file"""
    doc = Document()
    
    # Configure the entire document to automatically add space after paragraphs
    style = doc.styles['Normal']
    style.paragraph_format.space_after = Pt(12)
    
    doc.add_heading("Detailed Study Notes", 0)
    
    for line in request.notes_content.split('\n'):
        line = line.strip()
        
        if not line:
            continue # We don't need manual empty lines anymore!
            
        # Clean out Markdown bolding asterisks
        clean_line = line.replace('**', '')
            
        if line.startswith('### '):
            heading = doc.add_heading(clean_line.replace('### ', ''), 3)
            heading.paragraph_format.space_after = Pt(6)
        elif line.startswith('## '):
            heading = doc.add_heading(clean_line.replace('## ', ''), 2)
            heading.paragraph_format.space_after = Pt(6)
        elif line.startswith('# '):
            heading = doc.add_heading(clean_line.replace('# ', ''), 1)
            heading.paragraph_format.space_after = Pt(12)
        elif line.startswith('- ') or line.startswith('* '):
            doc.add_paragraph(clean_line[2:], style='List Bullet')
        else:
            doc.add_paragraph(clean_line)
            
    # Save to a temporary file
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(temp_file.name)
    
    return FileResponse(
        temp_file.name,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="Detailed_Study_Notes.docx"
    )

@app.post("/generate-question-bank")
def generate_question_bank(
    request: QuestionBankRequest, 
    current_user: models.User = Depends(get_current_user)
):
    """Generates a structured question bank with model answers"""
    agent = QuestionBankGenerator()
    q_bank = agent.generate_question_bank(request.document_name, request.topic)
    return {"question_bank": q_bank}

@app.get("/user-stats")
def get_user_stats(current_user: models.User = Depends(get_current_user)):
    """Returns the current user's usage statistics for the dashboard"""
    return {
        "documents_analyzed": current_user.documents_analyzed,
        "flashcards_generated": current_user.flashcards_generated,
        "study_guides_created": current_user.study_guides_created,
        "current_streak": current_user.current_streak # <-- ADDED STREAK TRACKING
    }